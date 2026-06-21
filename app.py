from functools import wraps
from io import BytesIO

from flask import (
    Flask, render_template, request, session,
    redirect, url_for, flash, send_file,
)
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

import db
from config import SECRET_KEY

app = Flask(__name__)
app.secret_key = SECRET_KEY


# -------------------------------------------------------------------
# Utilidades
# -------------------------------------------------------------------
def login_required(view):
    @wraps(view)
    def wrapper(*args, **kwargs):
        if "usuario" not in session:
            return redirect(url_for("login"))
        return view(*args, **kwargs)
    return wrapper


# -------------------------------------------------------------------
# Login / Logout
# -------------------------------------------------------------------
@app.route("/", methods=["GET", "POST"])
def login():
    if request.method == "POST":
        usuario = request.form.get("usuario", "").strip()
        password = request.form.get("password", "")

        # Comparacion en texto plano (decision del proyecto).
        user = db.fetch_one(
            """
            SELECT IdUsuario, Usuario, NombreCompleto, Rol
            FROM Usuarios
            WHERE Usuario = ? AND PasswordHash = ? AND Activo = 1
            """,
            (usuario, password),
        )

        if user:
            session["usuario"] = user["Usuario"]
            session["id_usuario"] = user["IdUsuario"]
            session["nombre"] = user["NombreCompleto"]
            session["rol"] = user["Rol"]
            return redirect(url_for("dashboard"))

        flash("Usuario o contrasena incorrectos.", "danger")

    return render_template("login.html")


@app.route("/logout")
def logout():
    session.clear()
    return redirect(url_for("login"))


# -------------------------------------------------------------------
# Dashboard
# -------------------------------------------------------------------
@app.route("/dashboard")
@login_required
def dashboard():
    kpi = db.fetch_one("SELECT * FROM vw_Dashboard")

    hoy = db.fetch_one(
        """
        SELECT
            ISNULL(SUM(CASE WHEN TipoMovimiento = 'E' THEN Cantidad END), 0) AS Entradas,
            ISNULL(SUM(CASE WHEN TipoMovimiento = 'S' THEN Cantidad END), 0) AS Salidas
        FROM MovimientosInventario
        WHERE CAST(FechaMovimiento AS DATE) = CAST(GETDATE() AS DATE)
        """
    )

    movimientos = db.fetch_all(
        """
        SELECT TOP 10
            FORMAT(m.FechaMovimiento, 'dd/MM/yyyy HH:mm') AS fecha,
            p.SKU AS sku,
            p.Nombre AS producto,
            CASE m.TipoMovimiento WHEN 'E' THEN 'Entrada' ELSE 'Salida' END AS tipo,
            m.Cantidad AS cantidad
        FROM MovimientosInventario m
        JOIN Productos p ON p.IdProducto = m.IdProducto
        ORDER BY m.FechaMovimiento DESC
        """
    )

    return render_template(
        "dashboard.html",
        total_productos=kpi["TotalProductos"],
        stock_bajo=kpi["ProductosCriticos"],
        entradas_hoy=hoy["Entradas"],
        salidas_hoy=hoy["Salidas"],
        movimientos=movimientos,
    )


# -------------------------------------------------------------------
# Usuario (perfil + cambio de contrasena)
# -------------------------------------------------------------------
@app.route("/usuarios")
@login_required
def usuarios():
    perfil = db.fetch_one(
        """
        SELECT Usuario, NombreCompleto, Rol,
               FORMAT(FechaRegistro, 'dd/MM/yyyy HH:mm') AS FechaRegistro,
               CASE WHEN Activo = 1 THEN 'Activo' ELSE 'Inactivo' END AS Estado
        FROM Usuarios
        WHERE IdUsuario = ?
        """,
        (session["id_usuario"],),
    )
    return render_template("usuarios.html", perfil=perfil)


@app.route("/usuarios/password", methods=["POST"])
@login_required
def cambiar_password():
    actual = request.form.get("actual", "")
    nueva = request.form.get("nueva", "")
    confirmar = request.form.get("confirmar", "")

    user = db.fetch_one(
        "SELECT PasswordHash FROM Usuarios WHERE IdUsuario = ?",
        (session["id_usuario"],),
    )

    if user["PasswordHash"] != actual:
        flash("La contrasena actual no es correcta.", "danger")
    elif not nueva:
        flash("La nueva contrasena no puede estar vacia.", "warning")
    elif nueva != confirmar:
        flash("La nueva contrasena y su confirmacion no coinciden.", "warning")
    else:
        db.execute(
            "UPDATE Usuarios SET PasswordHash = ? WHERE IdUsuario = ?",
            (nueva, session["id_usuario"]),
        )
        flash("Contrasena actualizada correctamente.", "success")

    return redirect(url_for("usuarios"))


# -------------------------------------------------------------------
# Kardex (detalle total de productos)
# -------------------------------------------------------------------
@app.route("/kardex")
@login_required
def kardex():
    productos = db.fetch_all(
        """
        SELECT
            p.SKU, p.Nombre, c.Nombre AS Categoria,
            p.Tamano, p.Color,
            p.StockActual, p.StockMinimo,
            p.PrecioCosto, p.PrecioVenta,
            CASE WHEN p.Activo = 1 THEN 'Activo' ELSE 'Inactivo' END AS Estado,
            CASE WHEN p.StockActual <= p.StockMinimo THEN 1 ELSE 0 END AS Critico
        FROM Productos p
        JOIN Categorias c ON c.IdCategoria = p.IdCategoria
        ORDER BY p.Nombre
        """
    )
    return render_template("kardex.html", productos=productos)


# -------------------------------------------------------------------
# Productos (busqueda, creacion, ingreso/salida)
# -------------------------------------------------------------------
def _consultar_productos(q):
    base = """
        SELECT
            p.IdProducto, p.SKU, p.Nombre, c.Nombre AS Categoria,
            p.StockActual, p.StockMinimo, p.PrecioVenta,
            CASE WHEN p.StockActual <= p.StockMinimo THEN 1 ELSE 0 END AS Critico
        FROM Productos p
        JOIN Categorias c ON c.IdCategoria = p.IdCategoria
    """
    if q:
        like = f"%{q}%"
        return db.fetch_all(
            base + " WHERE p.SKU LIKE ? OR p.Nombre LIKE ? OR c.Nombre LIKE ? ORDER BY p.Nombre",
            (like, like, like),
        )
    return db.fetch_all(base + " ORDER BY p.Nombre")


@app.route("/productos")
@login_required
def productos():
    q = request.args.get("q", "").strip()
    return render_template(
        "productos.html",
        productos=_consultar_productos(q),
        categorias=db.fetch_all("SELECT IdCategoria, Nombre FROM Categorias ORDER BY Nombre"),
        q=q,
    )


@app.route("/productos/crear", methods=["POST"])
@login_required
def crear_producto():
    f = request.form
    try:
        db.execute(
            """
            INSERT INTO Productos
                (SKU, Nombre, IdCategoria, PesoKg, Tamano, Color,
                 PrecioCosto, PrecioVenta, StockActual, StockMinimo)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                f.get("sku", "").strip(),
                f.get("nombre", "").strip(),
                int(f["categoria"]),
                float(f.get("peso") or 0),
                f.get("tamano", "").strip(),
                f.get("color", "").strip(),
                float(f.get("precio_costo") or 0),
                float(f.get("precio_venta") or 0),
                int(f.get("stock_inicial") or 0),
                int(f.get("stock_minimo") or 0),
            ),
        )
        flash(f"Producto '{f.get('nombre')}' creado correctamente.", "success")
    except Exception as e:
        flash(f"No se pudo crear el producto: {e}", "danger")
    return redirect(url_for("productos"))


@app.route("/productos/movimiento", methods=["POST"])
@login_required
def movimiento_producto():
    f = request.form
    try:
        id_producto = int(f["id_producto"])
        tipo = f.get("tipo")
        cantidad = int(f.get("cantidad") or 0)
        precio = float(f["precio"]) if f.get("precio") else None
        observacion = f.get("observacion", "").strip() or None

        if tipo not in ("E", "S"):
            raise ValueError("Tipo de movimiento invalido.")
        if cantidad <= 0:
            raise ValueError("La cantidad debe ser mayor a cero.")

        db.registrar_movimiento(
            id_producto, tipo, cantidad, precio, observacion, session["id_usuario"]
        )
        etiqueta = "Ingreso" if tipo == "E" else "Salida"
        flash(f"{etiqueta} registrado: {cantidad} unidad(es).", "success")
    except ValueError as e:
        flash(str(e), "warning")
    except Exception as e:
        flash(f"No se pudo registrar el movimiento: {e}", "danger")
    return redirect(url_for("productos", q=request.args.get("q", "")))


# -------------------------------------------------------------------
# Reportes (movimientos / stock + exportacion a Excel)
# -------------------------------------------------------------------
def _reporte_movimientos(f):
    where, params = [], []
    if f.get("desde"):
        where.append("CAST(m.FechaMovimiento AS DATE) >= ?")
        params.append(f["desde"])
    if f.get("hasta"):
        where.append("CAST(m.FechaMovimiento AS DATE) <= ?")
        params.append(f["hasta"])
    if f.get("categoria"):
        where.append("c.IdCategoria = ?")
        params.append(int(f["categoria"]))
    if f.get("producto"):
        where.append("p.IdProducto = ?")
        params.append(int(f["producto"]))
    if f.get("tipo") in ("E", "S"):
        where.append("m.TipoMovimiento = ?")
        params.append(f["tipo"])

    clause = ("WHERE " + " AND ".join(where)) if where else ""
    filas = db.fetch_all(
        f"""
        SELECT
            FORMAT(m.FechaMovimiento, 'dd/MM/yyyy HH:mm') AS Fecha,
            p.SKU, p.Nombre AS Producto, c.Nombre AS Categoria,
            CASE m.TipoMovimiento WHEN 'E' THEN 'Entrada' ELSE 'Salida' END AS Tipo,
            m.Cantidad, ISNULL(m.PrecioUnitario, 0) AS PrecioUnitario,
            ISNULL(u.NombreCompleto, '-') AS Usuario
        FROM MovimientosInventario m
        JOIN Productos p ON p.IdProducto = m.IdProducto
        JOIN Categorias c ON c.IdCategoria = p.IdCategoria
        LEFT JOIN Usuarios u ON u.IdUsuario = m.IdUsuario
        {clause}
        ORDER BY m.FechaMovimiento DESC
        """,
        params,
    )
    encabezados = ["Fecha", "SKU", "Producto", "Categoria", "Tipo",
                   "Cantidad", "Precio Unit.", "Usuario"]
    return "Reporte_Movimientos", encabezados, filas


def _reporte_stock(f):
    where, params = [], []
    if f.get("categoria"):
        where.append("c.IdCategoria = ?")
        params.append(int(f["categoria"]))
    if f.get("solo_bajo") == "1":
        where.append("p.StockActual <= p.StockMinimo")

    clause = ("WHERE " + " AND ".join(where)) if where else ""
    filas = db.fetch_all(
        f"""
        SELECT
            p.SKU, p.Nombre AS Producto, c.Nombre AS Categoria,
            p.StockActual, p.StockMinimo, p.PrecioVenta,
            CAST(p.StockActual * p.PrecioVenta AS DECIMAL(12,2)) AS ValorTotal
        FROM Productos p
        JOIN Categorias c ON c.IdCategoria = p.IdCategoria
        {clause}
        ORDER BY p.Nombre
        """,
        params,
    )
    encabezados = ["SKU", "Producto", "Categoria", "Stock", "Minimo",
                   "Precio Venta", "Valor Total"]
    return "Reporte_Stock", encabezados, filas


def _construir_reporte(f):
    if f.get("tipo_reporte") == "stock":
        return "stock", _reporte_stock(f)
    return "movimientos", _reporte_movimientos(f)


@app.route("/reportes", methods=["GET", "POST"])
@login_required
def reportes():
    contexto = {
        "categorias": db.fetch_all("SELECT IdCategoria, Nombre FROM Categorias ORDER BY Nombre"),
        "productos": db.fetch_all("SELECT IdProducto, Nombre FROM Productos ORDER BY Nombre"),
        "filtros": request.form,
        "tipo_reporte": request.form.get("tipo_reporte", "movimientos"),
        "encabezados": None,
        "filas": None,
    }

    if request.method == "POST":
        tipo_reporte, (_, encabezados, filas) = _construir_reporte(request.form)
        contexto["tipo_reporte"] = tipo_reporte
        contexto["encabezados"] = encabezados
        contexto["filas"] = filas

    return render_template("reportes.html", **contexto)


@app.route("/reportes/exportar", methods=["POST"])
@login_required
def exportar_reporte():
    _, (nombre, encabezados, filas) = _construir_reporte(request.form)

    wb = Workbook()
    ws = wb.active
    ws.title = nombre[:31]

    titulo = "Reporte de Movimientos" if nombre == "Reporte_Movimientos" else "Reporte de Stock"
    ws.append([f"Peluqueria Mary - {titulo}"])
    ws.append([])

    encabezado_fill = PatternFill("solid", fgColor="1976D2")
    encabezado_font = Font(bold=True, color="FFFFFF")
    ws.append(encabezados)
    for celda in ws[ws.max_row]:
        celda.fill = encabezado_fill
        celda.font = encabezado_font
        celda.alignment = Alignment(horizontal="center")

    for fila in filas:
        ws.append(list(fila.values()))

    for i, _ in enumerate(encabezados, start=1):
        ws.column_dimensions[ws.cell(row=3, column=i).column_letter].width = 20

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return send_file(
        buffer,
        as_attachment=True,
        download_name=f"{nombre}.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


if __name__ == "__main__":
    app.run(debug=True)
